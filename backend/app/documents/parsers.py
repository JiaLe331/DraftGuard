from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader

MAX_BYTES = 10 * 1024 * 1024
MAX_PAGES = 20
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_CELLS = 100_000
MAX_TEXT = 1_000_000
EXTENSIONS = {".txt", ".pdf", ".docx", ".xlsx"}


class DocumentProblem(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def parse_document(path: Path, filename: str, document_id: str) -> dict:
    """Source units are also the text shown in the evidence viewer."""
    units: list[dict] = []
    size = 0

    def add(text: str, locator: str, **location):
        nonlocal size
        size += len(text)
        if size > MAX_TEXT:
            raise DocumentProblem("limit_exceeded", "Extracted text exceeds the supported limit.")
        if text.strip():
            units.append(
                {
                    "id": f"u{len(units) + 1}",
                    "document_id": document_id,
                    "locator": locator,
                    "text": text,
                    **location,
                }
            )

    ext = Path(filename).suffix.lower()
    try:
        if ext not in EXTENSIONS:
            raise DocumentProblem("unsupported_format", "This file format is not supported.")
        if path.stat().st_size > MAX_BYTES:
            raise DocumentProblem("limit_exceeded", "Files must be no larger than 10 MB.")
        if not path.stat().st_size:
            raise DocumentProblem("empty_file", "The document is empty. Supply a readable source.")
        if ext == ".txt":
            text = path.read_text(encoding="utf-8-sig")
            if "\x00" in text or any(ord(c) < 32 and c not in "\n\r\t" for c in text):
                raise DocumentProblem("invalid_format", "The file is not supported plain text.")
            for line, text_line in enumerate(text.splitlines(), 1):
                add(text_line, f"Line {line}", line=line)
        elif ext == ".pdf":
            with path.open("rb") as handle:
                if not handle.read(1024).lstrip().startswith(b"%PDF-"):
                    raise DocumentProblem("invalid_format", "The file is not a readable PDF.")
                handle.seek(0)
                reader = PdfReader(handle)
                if reader.is_encrypted:
                    raise DocumentProblem("encrypted", "Encrypted PDFs must be replaced.")
                if len(reader.pages) > MAX_PAGES:
                    raise DocumentProblem("limit_exceeded", "PDFs must have at most 20 pages.")
                for page_number, page in enumerate(reader.pages, 1):
                    text = (
                        (page.extract_text(extraction_mode="layout") or "")
                        if page.get("/Contents")
                        else ""
                    )
                    for line, text_line in enumerate(text.splitlines(), 1):
                        add(
                            text_line,
                            f"Page {page_number}, line {line}",
                            page=page_number,
                            line=line,
                        )
        else:
            with ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
                    raise DocumentProblem("limit_exceeded", "Document archive exceeds safe limits.")
                expected = "word/document.xml" if ext == ".docx" else "xl/workbook.xml"
                if expected not in archive.namelist():
                    raise DocumentProblem(
                        "invalid_format", "Document content does not match its type."
                    )
                if any(e.flag_bits & 1 for e in entries):
                    raise DocumentProblem("encrypted", "Encrypted archives are not supported.")
            if ext == ".docx":
                document = Document(path)
                for block_index, block in enumerate(document.iter_inner_content(), 1):
                    if isinstance(block, Paragraph):
                        add(block.text, f"Block {block_index}, paragraph", block=block_index)
                    elif isinstance(block, Table):
                        for row_index, row in enumerate(block.rows, 1):
                            cells = []
                            seen = set()
                            for cell in row.cells:
                                if cell._tc not in seen:
                                    seen.add(cell._tc)
                                    cells.append(cell.text)
                            add(
                                " | ".join(cells),
                                f"Table at block {block_index}, row {row_index}, "
                                f"cells 1–{len(cells)}",
                                block=block_index,
                                row=row_index,
                            )
            else:
                book = load_workbook(
                    BytesIO(path.read_bytes()), read_only=True, data_only=False, keep_links=False
                )
                cached = None
                try:
                    cached = load_workbook(
                        BytesIO(path.read_bytes()), read_only=True, data_only=True, keep_links=False
                    )
                    if len(book.worksheets) > 20:
                        raise DocumentProblem("limit_exceeded", "Workbook exceeds 20 sheets.")
                    visited = 0
                    for sheet, values_sheet in zip(book.worksheets, cached.worksheets, strict=True):
                        if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_CELLS:
                            raise DocumentProblem(
                                "limit_exceeded", "Workbook exceeds 100,000 cells."
                            )
                        sheet.reset_dimensions()
                        values_sheet.reset_dimensions()
                        for row, values in zip(sheet.rows, values_sheet.rows, strict=True):
                            visited += len(row)
                            if visited > MAX_CELLS:
                                raise DocumentProblem(
                                    "limit_exceeded", "Workbook exceeds 100,000 cells."
                                )
                            cells = []
                            coordinates = []
                            for cell, cached_cell in zip(row, values, strict=True):
                                if cell.value is None:
                                    continue
                                value = cell.value
                                if cell.data_type == "f":
                                    value = cached_cell.value
                                    if value is None:
                                        value = "[Formula result unavailable]"
                                cells.append(str(value))
                                coordinates.append(cell.coordinate)
                            if coordinates:
                                add(
                                    " | ".join(cells),
                                    f"{sheet.title}!{','.join(coordinates)}",
                                    sheet=sheet.title,
                                    cells=coordinates,
                                )
                finally:
                    book.close()
                    if cached:
                        cached.close()
        if not units:
            code = "visual_extraction_required" if ext == ".pdf" else "no_text"
            message = (
                "Visual extraction required. This PDF has no usable text layer."
                if ext == ".pdf"
                else "The document has no usable text."
            )
            raise DocumentProblem(code, message)
        return {"state": "PARSED", "units": units, "error": None}
    except DocumentProblem as exc:
        return {
            "state": "UNREADABLE",
            "units": [],
            "error": {"code": exc.code, "message": str(exc)},
        }
    except (OSError, UnicodeError, BadZipFile):
        return {
            "state": "UNREADABLE",
            "units": [],
            "error": {"code": "unreadable", "message": "The source could not be read."},
        }
    except Exception:
        # Parser-specific errors must not expose document bodies or local paths.
        return {
            "state": "UNREADABLE",
            "units": [],
            "error": {"code": "malformed_document", "message": "The document is malformed."},
        }
