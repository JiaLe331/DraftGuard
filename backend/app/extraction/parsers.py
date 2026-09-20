"""Bounded document readers that retain displayable source locations."""

import codecs
import re
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from defusedxml.ElementTree import fromstring
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from pypdf import PdfReader

from .models import ExtractionLimits, Format, SourceUnit


class ParseProblem(Exception):
    def __init__(self, code: str, message: str, action: str, *, rejected: bool = False):
        self.code, self.message, self.action, self.rejected = code, message, action, rejected
        super().__init__(message)


def limit_problem(message: str) -> ParseProblem:
    return ParseProblem("RESOURCE_LIMIT", message, "Upload a smaller document.", rejected=True)


class Units:
    def __init__(self, document_id: str, limits: ExtractionLimits):
        self.document_id, self.limits = document_id, limits
        self.items: list[SourceUnit] = []
        self.characters = 0

    def add(self, text: str, **locator):
        self.characters += len(text)
        if self.characters > self.limits.max_text_chars:
            raise limit_problem("The document contains too much extracted text.")
        if len(self.items) >= self.limits.max_source_units:
            raise limit_problem("The document contains too many source units.")
        self.items.append(
            SourceUnit(
                unit_id=f"unit-{len(self.items) + 1}",
                document_id=self.document_id,
                text=text,
                **locator,
            )
        )


def detect_format(content: bytes, filename: str, limits: ExtractionLimits) -> Format:
    if len(content) > limits.max_file_bytes:
        raise limit_problem("The file exceeds the configured byte limit.")
    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix not in {"txt", "pdf", "docx", "xlsx"}:
        raise ParseProblem(
            "UNSUPPORTED_FORMAT",
            "Only TXT, PDF, DOCX, and XLSX are supported.",
            "Provide a supported document format.",
            rejected=True,
        )
    is_pdf = content.lstrip().startswith(b"%PDF-")
    is_zip = content.startswith(b"PK")
    is_ole = content.startswith(bytes.fromhex("d0cf11e0a1b11ae1"))
    if is_ole and suffix in {"docx", "xlsx"}:
        raise ParseProblem(
            "ENCRYPTED_OR_LEGACY_OFFICE",
            "Encrypted or legacy Office input is unsupported.",
            "Export an unencrypted DOCX or XLSX file.",
            rejected=True,
        )
    if (
        (suffix == "pdf" and not is_pdf)
        or (suffix in {"docx", "xlsx"} and not is_zip)
        or (suffix == "txt" and (is_pdf or is_zip or is_ole))
    ):
        raise ParseProblem(
            "FORMAT_MISMATCH",
            "The file contents do not match its extension.",
            "Export the document in its correct format.",
            rejected=True,
        )
    return suffix


def _validate_office(content: bytes, kind: Format, limits: ExtractionLimits):
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > limits.max_archive_entries:
                raise limit_problem("The Office archive contains too many entries.")
            if sum(item.file_size for item in entries) > limits.max_archive_bytes:
                raise limit_problem("The expanded Office archive is too large.")
            if any(item.flag_bits & 1 for item in entries):
                raise ParseProblem(
                    "ENCRYPTED_DOCUMENT",
                    "Encrypted archives are unsupported.",
                    "Provide an unencrypted document.",
                    rejected=True,
                )
            names = {item.filename for item in entries}
            required = "word/document.xml" if kind == "docx" else "xl/workbook.xml"
            other = "xl/workbook.xml" if kind == "docx" else "word/document.xml"
            if required not in names or other in names or "[Content_Types].xml" not in names:
                raise ParseProblem(
                    "FORMAT_MISMATCH",
                    "The archive is not the declared Office format.",
                    "Export a valid DOCX or XLSX document.",
                    rejected=True,
                )
            if len(names) != len(entries):
                raise ValueError("Duplicate archive entries")
            if any(name.lower().endswith("vbaproject.bin") for name in names):
                raise ParseProblem(
                    "UNSUPPORTED_FORMAT",
                    "Macro-enabled Office files are unsupported.",
                    "Export a document without macros.",
                    rejected=True,
                )
            # Reject dangerous XML before either Office library handles the package.
            for item in entries:
                if item.filename.endswith((".xml", ".rels")):
                    fromstring(archive.read(item), forbid_dtd=True)
            if kind == "xlsx":
                _validate_worksheets(archive, limits)
    except ParseProblem:
        raise
    except (BadZipFile, ValueError, KeyError, RuntimeError) as exc:
        raise ParseProblem(
            "MALFORMED_DOCUMENT",
            "The Office archive could not be read.",
            "Export a new copy of the document.",
        ) from exc


def _validate_worksheets(archive: ZipFile, limits: ExtractionLimits):
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    workbook = fromstring(archive.read("xl/workbook.xml"), forbid_dtd=True)
    if len(workbook.findall(f"{ns}sheets/{ns}sheet")) > limits.max_sheets:
        raise limit_problem("The workbook contains too many sheets.")
    cell_budget = 0
    sheets = [n for n in archive.namelist() if re.fullmatch(r"xl/worksheets/[^/]+\.xml", n)]
    if len(sheets) > limits.max_sheets:
        raise limit_problem("The workbook contains too many worksheet parts.")
    for name in sheets:
        tree = fromstring(archive.read(name), forbid_dtd=True)
        max_row, max_col = 0, 0
        dimension = tree.find(f"{ns}dimension")
        if dimension is not None:
            _, _, max_col, max_row = range_boundaries(dimension.attrib["ref"])
            max_row, max_col = max_row or 0, max_col or 0
        # Inspect actual coordinates too: a forged A1:A1 dimension must not hide later cells.
        for row_node in tree.iter(f"{ns}row"):
            max_row = max(max_row, int(row_node.attrib.get("r", "0")))
        for cell in tree.iter(f"{ns}c"):
            row, col = coordinate_to_tuple(cell.attrib["r"])
            max_row, max_col = max(max_row, row), max(max_col, col)
        if max_row > limits.max_rows or max_col > limits.max_columns:
            raise limit_problem("The workbook exceeds the configured row or column limit.")
        cell_budget += max_row * max_col
        if cell_budget > limits.max_cells:
            raise limit_problem("The workbook exceeds the inspected-cell limit.")


def _txt(content: bytes, units: Units):
    try:
        encoding = (
            "utf-16"
            if content.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE))
            else "utf-8-sig"
        )
        text = content.decode(encoding, errors="strict")
    except UnicodeError as exc:
        raise ParseProblem(
            "UNSUPPORTED_ENCODING",
            "The text is not valid UTF-8 or BOM-marked UTF-16.",
            "Save the text file as UTF-8.",
            rejected=True,
        ) from exc
    if any(ord(c) < 32 and c not in "\n\r\t\f" for c in text):
        raise ParseProblem(
            "BINARY_TEXT",
            "The text file contains binary control characters.",
            "Provide a plain-text document.",
            rejected=True,
        )
    if re.match(r"\s*(?:<!doctype\s+html|<html\b|<script\b|<\?xml)", text, re.I):
        raise ParseProblem(
            "FORMAT_MISMATCH",
            "Markup supplied as a text document is unsupported.",
            "Export the source as plain text.",
            rejected=True,
        )
    for line, text_line in enumerate(text.splitlines(), 1):
        units.add(text_line, line=line)


def _pdf(content: bytes, units: Units):
    reader = PdfReader(BytesIO(content), strict=True)
    if reader.is_encrypted:
        raise ParseProblem(
            "ENCRYPTED_DOCUMENT",
            "Encrypted PDFs are unsupported.",
            "Provide an unencrypted PDF.",
            rejected=True,
        )
    if len(reader.pages) > units.limits.max_pdf_pages:
        raise limit_problem("The PDF exceeds the configured page limit.")
    for page_number, page in enumerate(reader.pages, 1):
        stream = page.get_contents()
        if stream is not None and len(stream.get_data()) > units.limits.max_pdf_stream_bytes:
            raise limit_problem("A PDF page contains too much decompressed content.")
        text = page.extract_text() or ""
        units.add(text, page=page_number)


def _docx(content: bytes, units: Units):
    document = Document(BytesIO(content))
    paragraph_index, table_index = 0, 0

    def read_table(table: Table):
        nonlocal table_index
        table_index += 1
        current_table = table_index
        seen = set()
        for row_number, row in enumerate(table.rows, 1):
            for column_number, cell in enumerate(row.cells, 1):
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                units.add(cell.text, table=current_table, row=row_number, column=column_number)
                for nested in cell.tables:
                    read_table(nested)

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_index += 1
            units.add(block.text, paragraph=paragraph_index)
        else:
            read_table(block)


def _xlsx(content: bytes, units: Units):
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    inspected = 0
    try:
        for sheet in workbook.worksheets:
            # Preflight already checked both declared and actual coordinates.
            sheet.reset_dimensions()
            for row in sheet.iter_rows():
                inspected += len(row)
                if inspected > units.limits.max_cells:
                    raise limit_problem("The workbook exceeds the inspected-cell limit.")
                for cell in row:
                    if cell.value is None:
                        continue
                    if cell.row > units.limits.max_rows or cell.column > units.limits.max_columns:
                        raise limit_problem("The workbook exceeds its dimension limits.")
                    units.add(
                        str(cell.value),
                        sheet=sheet.title,
                        cell=cell.coordinate,
                        row=cell.row,
                        column=cell.column,
                        is_formula=cell.data_type == "f",
                    )
    finally:
        workbook.close()


def parse(content: bytes, kind: Format, document_id: str, limits: ExtractionLimits):
    units = Units(document_id, limits)
    if kind in {"docx", "xlsx"}:
        _validate_office(content, kind, limits)
    {"txt": _txt, "pdf": _pdf, "docx": _docx, "xlsx": _xlsx}[kind](content, units)
    return units.items
