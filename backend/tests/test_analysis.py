import json
import os
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.documents.analysis import FIELDS, adapt_document, analyze, classify
from app.extraction import extract_document
from app.extraction.models import ExtractionLimits
from app.store import Store


def make_pdf(path, text):
    writer = PdfWriter()
    page = writer.add_blank_page(600, 800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): writer._add_object(font),
                }
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 40 740 Td ({text}) Tj ET".encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


@pytest.mark.parametrize(
    ("subject", "body", "category"),
    [
        ("Check draft BL", "Please compare the SI and draft BL.", "BL_COMPARISON"),
        (
            "REQUEST SI",
            "Please find Shipping instruction.\nPlease revert with draft BL once available.",
            "SI_REQUEST",
        ),
        ("Billing query", "Please cancel invoice 123.", "INVOICE_QUERY"),
        ("Holiday announcement", "Office resumes tomorrow.", "GENERAL"),
        ("You have won a gift card", "Claim your prize!", "SPAM"),
        ("Unclear request", "Can you help with this?", None),
        ("REQUEST SI", "Please compare the SI and draft BL.", None),
        (
            "Holiday",
            "Office resumes tomorrow.\nFrom: quoted@example.test\nCheck draft BL",
            "GENERAL",
        ),
    ],
)
def test_classification_uses_primary_intent(subject, body, category):
    assert classify(subject, body)["category"] == category


def parsed(path, filename=None, limits=None):
    result = extract_document(path.read_bytes(), filename or path.name, "version-id", limits=limits)
    return adapt_document({"document_id": "version-id", "result": result.model_dump(mode="json")})


def test_four_formats_keep_locators_and_formula_uncertainty(tmp_path):
    txt = tmp_path / "source.txt"
    txt.write_text("SHIPPING INSTRUCTION\nGross Weight (KG): 12,345 KG")
    pdf = tmp_path / "source.pdf"
    make_pdf(pdf, "Gross Weight: 12,345 KG")
    docx = tmp_path / "source.docx"
    word = Document()
    word.add_paragraph("SHIPPING INSTRUCTION")
    cells = word.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text, cells[1].text = "Gross Weight (KG)", "12,345 KG"
    word.save(docx)
    xlsx = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.append(["Gross Weight (KG)", "12,345 KG"])
    book.save(xlsx)
    for path, location in [(txt, "Line"), (pdf, "Page"), (docx, "Table"), (xlsx, "Sheet!")]:
        document = parsed(path)
        value = document["values"]["gross_weight_kg"]
        assert value["normalized_value"] == "12345"
        evidence = value["evidence"][0]
        assert location in evidence["locator"] and evidence["document_id"] == "version-id"
        assert any(evidence["excerpt"] in unit["text"] for unit in document["units"])
    book.active["B1"] = "=12000+345"
    book.save(xlsx)
    field = parsed(xlsx)["values"]["gross_weight_kg"]
    assert field["normalized_value"] is None
    assert field["reason"] == "FORMULA_VALUE"


def test_scans_empty_corrupt_encrypted_and_spoofed_documents(tmp_path):
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(600, 800)
    writer.write(path)
    assert parsed(path, "scan.pdf")["error"]["code"] == "VISUAL_REVIEW_REQUIRED"
    writer.encrypt("secret")
    writer.write(path)
    assert parsed(path, "scan.pdf")["error"]["code"] == "ENCRYPTED_DOCUMENT"
    path.write_bytes(b"%PDF-broken")
    assert parsed(path, "scan.pdf")["error"]["code"] == "MALFORMED_DOCUMENT"
    path.write_bytes(b"")
    assert parsed(path, "scan.pdf")["state"] == "UNREADABLE"
    path.write_bytes(b"this is not an xlsx")
    assert parsed(path, "source.xlsx")["state"] == "UNREADABLE"


def test_comparison_preserves_shared_port_values_and_review_states(tmp_path):
    first = tmp_path / "si.txt"
    first.write_text("SHIPPING INSTRUCTION\nPOL: NHAVA SHEVA, INDIA (INNSA)\nGross Weight: 1234")
    second = tmp_path / "bl.txt"
    second.write_text("BILL OF LADING\nPOL: NHAVA SHEVA, INDIA\nGross Weight: 1234 KG")
    result = analyze(
        {"subject": "Check draft BL", "body": ""},
        [
            {"id": str(index), "path": str(path), "filename": path.name}
            for index, path in enumerate((first, second))
        ],
    )
    fields = {f["key"]: f for f in result["fields"]}
    assert fields["port_of_loading"]["finding"] == "MISMATCH"
    assert fields["port_of_loading"]["si"]["normalized_value"] == "NHAVA SHEVA, INDIA (INNSA)"
    assert fields["gross_weight_kg"]["finding"] == "NEEDS_REVIEW"
    assert fields["gross_weight_kg"]["si"]["raw_value"] == "1234"


def test_missing_attachments_never_match():
    result = analyze(
        {"subject": "Check draft BL", "body": "Please compare the SI and draft BL"}, []
    )
    assert result["workflow_state"] == "WAITING_DOCUMENT"
    assert len(result["fields"]) == 7
    assert all(f["finding"] == "NOT_CHECKED" for f in result["fields"])


def test_provided_dataset_acceptance(tmp_path):
    default = (
        Path(__file__).resolve().parents[3] / "problem-statement/sdoc-hackathon-docker/data_v2"
    )
    root = Path(os.environ.get("DATASET_DIR", default))
    if not root.is_dir():
        pytest.skip("Set DATASET_DIR to run organizer-dataset integration checks.")
    store = Store(tmp_path / "local")
    assert len(store.import_dataset(root)) == 520
    assert store.import_dataset(root) == []
    listing = store.list_samples()
    assert listing["total"] == 520 and len(listing["items"]) == 50
    assert listing["summary"]["attachments"] == 250
    results = {}
    for email_id in ("email_004", "email_160", "email_055", "email_501", "email_516"):
        results[email_id] = store.analyze(email_id, 1)["current_run"]["result"]
    scan = store.analyze("email_512", 1)
    assert scan["current_run"] is None
    assert scan["latest_run"]["error"]["code"] == "AI_NOT_CONFIGURED"
    assert results["email_004"]["known_defect_fields"] == ["consignee", "notify_party"]
    si = results["email_160"]["fields"]
    assert {f["key"] for f in si} == set(FIELDS)
    assert all(
        f["si"]["normalized_value"] is not None and f["si"]["evidence"][0]["page"] == 1 for f in si
    )
    assert si[-1]["si"]["normalized_value"] == "23702"
    spreadsheet = results["email_055"]["fields"]
    assert all(f["finding"] == "MATCH" for f in spreadsheet[:6])
    assert spreadsheet[-1]["finding"] == "NEEDS_REVIEW"
    assert spreadsheet[-1]["si"]["reason"] == "MISSING_WEIGHT_UNIT"
    assert spreadsheet[-1]["bl"]["normalized_value"] == "243588"
    assert any(
        r["code"] == "WRONG_DOCUMENT_TYPE" for r in results["email_501"]["review_requirements"]
    )
    weight = results["email_516"]["fields"][-1]
    assert weight["si"]["normalized_value"] is None and weight["bl"]["normalized_value"] == "235550"
    # Results contain actual source excerpts, never references to the answer file.
    assert "ground_truth" not in json.dumps(results)


def test_document_limits_and_ambiguous_pairs(tmp_path):
    path = tmp_path / "large.txt"
    path.write_text("SHIPPING INSTRUCTION\nShipper: A LTD")
    assert (
        parsed(path, limits=ExtractionLimits(max_file_bytes=5))["error"]["code"] == "RESOURCE_LIMIT"
    )
    pdf = tmp_path / "large.pdf"
    writer = PdfWriter()
    for _ in range(21):
        writer.add_blank_page(600, 800)
    writer.write(pdf)
    assert parsed(pdf)["error"]["code"] == "RESOURCE_LIMIT"
    documents = [{"id": name, "filename": path.name, "path": str(path)} for name in ["one", "two"]]
    result = analyze({"subject": "Compare draft BL", "body": "Check SI and draft BL."}, documents)
    assert result["workflow_state"] == "REVIEW_REQUIRED"
    assert any(r["code"] == "ambiguous_document_pair" for r in result["review_requirements"])
    assert result["coverage"]["checked"] == 0
