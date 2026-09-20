"""Behavior tests use source documents, never organizer ground truth."""

import json
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.extraction import DocumentExtraction, ExtractionLimits, extract_document
from app.extraction.models import FIELD_KEYS, Locator

FIXTURES = Path(__file__).parent / "fixtures" / "extraction"
TEXT = """SHIPPING INSTRUCTION
Shipper: Alpha Paper PTE LTD
Consignee: Beta Trading LLC
Notify Party: Beta Trading LLC
POL: PORT KLANG (WESTPORT), MALAYSIA
POD: SINGAPORE
Container Count: 6 x 40'HC
Gross Weight (KG): 131,058 KG
"""


def extract_text(text=TEXT, **kwargs):
    return extract_document(text.encode(), "source.txt", "source-v1", **kwargs)


def field(result, key):
    return next(value for value in result.fields if value.field == key)


def codes(result):
    return {issue.code for issue in result.issues}


def fixture(name, role=None):
    return extract_document((FIXTURES / name).read_bytes(), name, f"fixture-{name}", role)


def assert_evidence(result):
    units = {unit.unit_id: unit for unit in result.source_units}
    assert len(units) == len(result.source_units)
    assert [value.field for value in result.fields] == list(FIELD_KEYS)
    for value in result.fields:
        if value.value_state == "PRESENT":
            assert value.normalized_value is not None
            assert value.evidence
        for evidence in value.evidence:
            source = units[evidence.unit_id]
            assert evidence.document_id == source.document_id == result.document_id
            assert evidence.verified
            assert evidence.excerpt == source.text
            for attribute in Locator.model_fields:
                assert getattr(evidence, attribute) == getattr(source, attribute)


def pdf_bytes(text="", pages=1, encrypted=False):
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    for _ in range(pages):
        page = writer.add_blank_page(width=612, height=792)
        if text:
            page[NameObject("/Resources")] = DictionaryObject(
                {
                    NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
                }
            )
            stream = DecodedStreamObject()
            lines = [
                line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
                for line in text.splitlines()
            ]
            stream.set_data(
                (
                    "BT /F1 12 Tf 14 TL 40 740 Td "
                    + " ".join(f"({line}) Tj T*" for line in lines)
                    + " ET"
                ).encode()
            )
            page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("test-password")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def workbook_bytes(rows=None, sheets=1):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Source"
    sheet.append(["SHIPPING INSTRUCTION"])
    for row in rows or []:
        sheet.append(row)
    for index in range(1, sheets):
        workbook.create_sheet(f"Sheet{index}")
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def docx_bytes():
    document = Document()
    document.add_paragraph("SHIPPING INSTRUCTION")
    document.add_paragraph("Shipper: Alpha LTD")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Gross Weight (KG)"
    table.cell(0, 1).text = "123.40"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def replace_zip_part(content, part, transform):
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            target.writestr(entry.filename, transform(data) if entry.filename == part else data)
    return output.getvalue()


def test_real_pdf_seven_fields_and_total_weight():
    result = fixture("email_160_SI.pdf", "SI")
    assert result.parsing_status == "READABLE"
    assert result.detected_role == "SI"
    assert not result.needs_review
    assert field(result, "gross_weight_kg").normalized_value == "23702"
    assert field(result, "container_count").normalized_value == "1"
    assert "ON BEHALF OF VITAL SOLUTIONS PTE LTD" in field(result, "shipper").normalized_value
    assert "77 ROBINSON ROAD" in field(result, "shipper").raw_value
    assert all(e.page == 1 for f in result.fields for e in f.evidence)
    assert_evidence(result)


def test_real_txt_values_preserve_company_discrepancies():
    si, bl = fixture("email_004_SI.txt"), fixture("email_004_BL.txt")
    for key in ["consignee", "notify_party"]:
        assert field(si, key).normalized_value == "EAST BRIGHT FZ-LLC"
        assert field(bl, key).normalized_value == "UAB NOVAKOPA"
    for result in [si, bl]:
        assert field(result, "gross_weight_kg").normalized_value == "131058"
        assert not result.needs_review
        assert all(e.line is not None for f in result.fields for e in f.evidence)
        assert_evidence(result)


def test_real_office_values_and_absent_weight_unit():
    docx, xlsx = fixture("email_055_BL.docx"), fixture("email_055_SI.xlsx")
    assert field(docx, "gross_weight_kg").raw_value == "243,588"
    assert field(docx, "gross_weight_kg").normalized_value == "243588"
    assert not docx.needs_review
    assert field(xlsx, "gross_weight_kg").raw_value == "243588"
    assert field(xlsx, "gross_weight_kg").normalized_value is None
    assert field(xlsx, "gross_weight_kg").value_state == "AMBIGUOUS"
    assert "MISSING_WEIGHT_UNIT" in codes(xlsx)
    assert {(e.sheet, e.cell) for e in field(xlsx, "gross_weight_kg").evidence} == {
        ("S.I.", "A10"),
        ("S.I.", "B10"),
    }
    assert all(e.table == 1 and e.row == 7 for e in field(docx, "gross_weight_kg").evidence)
    assert_evidence(docx)
    assert_evidence(xlsx)


def test_real_missing_weight_is_not_copied_from_bl():
    si, bl = fixture("email_516_SI.txt"), fixture("email_516_BL.txt")
    assert field(si, "gross_weight_kg").raw_value == "N/A"
    assert field(si, "gross_weight_kg").normalized_value is None
    assert field(si, "gross_weight_kg").value_state == "MISSING"
    assert field(bl, "gross_weight_kg").normalized_value == "235550"
    assert_evidence(si)


def test_real_image_pdf_is_not_a_successful_extraction():
    result = fixture("email_512_SI.pdf")
    assert result.parsing_status == "NO_USABLE_TEXT"
    assert result.needs_review
    assert "VISUAL_REVIEW_REQUIRED" in codes(result)
    assert all(f.normalized_value is None for f in result.fields)


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "utf-16"])
def test_supported_text_encodings(encoding):
    result = extract_document(TEXT.encode(encoding), "source.TXT", "version")
    assert not result.needs_review
    assert_evidence(result)


@pytest.mark.parametrize("value", ["N/A", "TBA", "______", "", "-", "NULL"])
def test_missing_values_are_not_zero(value):
    result = extract_text(TEXT.replace("131,058 KG", value))
    weight = field(result, "gross_weight_kg")
    assert weight.value_state == "MISSING"
    assert weight.normalized_value is None
    assert weight.raw_value == (value or None)


@pytest.mark.parametrize(
    ("label", "value", "expected"),
    [
        ("Gross Weight (KG)", "243,588", "243588"),
        ("Gross Wt (kgs)", "243588.000", "243588"),
        ("Gross Weight", "1.234 MT", "1234"),
        ("Gross Weight (grams)", "12345", "12.345"),
        ("Gross Weight", "0 KG", "0"),
        ("Gross Weight (metric tonnes)", "2.000001", "2000.001"),
    ],
)
def test_exact_weight_normalization(label, value, expected):
    result = extract_text(f"SHIPPING INSTRUCTION\n{label}: {value}\n")
    assert field(result, "gross_weight_kg").normalized_value == expected
    assert_evidence(result)


@pytest.mark.parametrize("value", ["12,50 KG", "1.234,56 KG", "-5 KG", "12 pounds", "1e5 KG"])
def test_ambiguous_or_unsupported_weights_need_review(value):
    result = extract_text(TEXT.replace("131,058 KG", value))
    assert field(result, "gross_weight_kg").value_state == "AMBIGUOUS"
    assert field(result, "gross_weight_kg").normalized_value is None


def test_conflicting_units_are_not_silently_converted():
    result = extract_text(TEXT.replace("131,058 KG", "2 MT"))
    assert "CONFLICTING_WEIGHT_UNITS" in codes(result)


@pytest.mark.parametrize(
    ("label", "value", "expected"),
    [
        ("No. of Containers or Packages", "6 x 40'HC", "6"),
        ("Container Count", "0", "0"),
        ("Total Containers", "12 × 20'FCL", "12"),
        ("No. of Containers or Packages", "45", None),
        ("Container Count", "ABCD1234567", None),
        ("Container Count", "100 PACKAGES", None),
        ("Container Count", "-1", None),
    ],
)
def test_container_quantity_not_ids_or_packages(label, value, expected):
    result = extract_text(f"SHIPPING INSTRUCTION\n{label}: {value}\n")
    assert field(result, "container_count").normalized_value == expected


def test_total_weight_wins_and_conflicting_totals_stay_ambiguous():
    text = "SHIPPING INSTRUCTION\nGross Weight (KG): 12\nTOTAL Gross Weight (KG): 30\n"
    assert field(extract_text(text), "gross_weight_kg").normalized_value == "30"
    conflict = extract_text(text + "TOTAL Gross Weight (KG): 31\n")
    assert field(conflict, "gross_weight_kg").value_state == "AMBIGUOUS"
    assert "CONFLICTING_VALUES" in codes(conflict)
    assert_evidence(conflict)


def test_conflicting_company_values_preserve_both_sources():
    result = extract_text(TEXT + "Consignee: Other Company LTD\n")
    assert field(result, "consignee").value_state == "AMBIGUOUS"
    assert "Other Company LTD" in field(result, "consignee").raw_value
    assert_evidence(result)


def test_split_labels_and_value_beneath_label():
    text = TEXT.replace(
        "POL: PORT KLANG (WESTPORT), MALAYSIA", "Port of\nLoading\nPORT KLANG (WESTPORT), MALAYSIA"
    )
    text = text.replace("Gross Weight (KG): 131,058 KG", "Gross\nWeight (KG)\n131,058 KG")
    result = extract_text(text)
    assert field(result, "port_of_loading").normalized_value == "PORT KLANG (WESTPORT), MALAYSIA"
    assert field(result, "gross_weight_kg").normalized_value == "131058"
    assert_evidence(result)


def test_agency_clause_is_retained_and_unclear_identity_is_reviewed():
    result = extract_text(
        TEXT.replace("Alpha Paper PTE LTD", "Alpha LTD\nON BEHALF OF Beta LTD\n12 Main Road")
    )
    assert field(result, "shipper").normalized_value == "ALPHA LTD ON BEHALF OF BETA LTD"
    unclear = extract_text(TEXT.replace("Alpha Paper PTE LTD", "Alpha LTD\nOther Company LLC"))
    assert field(unclear, "shipper").value_state == "AMBIGUOUS"
    assert "Other Company LLC" in field(unclear, "shipper").raw_value


@pytest.mark.parametrize(
    "continuation",
    [
        "ROAD RUNNER TRADING",
        "12 Main Road\nON BEHALF OF Beta LTD",
    ],
)
def test_additional_company_text_is_not_discarded_as_address(continuation):
    result = extract_text(TEXT.replace("Alpha Paper PTE LTD", f"Alpha LTD\n{continuation}"))
    assert field(result, "shipper").value_state == "AMBIGUOUS"
    assert continuation in field(result, "shipper").raw_value


def test_two_port_lines_require_review():
    result = extract_text(TEXT.replace("POD: SINGAPORE", "POD: SINGAPORE\nKARACHI"))
    assert field(result, "port_of_discharge").normalized_value is None
    assert "AMBIGUOUS_TEXT_VALUE" in codes(result)


@pytest.mark.parametrize(
    ("heading", "expected_role", "role", "code"),
    [
        ("BILL OF LADING INSTRUCTION", "SI", "SI", None),
        ("BL INSTRUCTION", "SI", "SI", None),
        ("BILL OF LADING (DRAFT)", "SI", "BL", "DOCUMENT_ROLE_MISMATCH"),
        ("COMMERCIAL INVOICE", "BL", None, "WRONG_DOCUMENT_TYPE"),
        ("UNLABELED DOCUMENT", "BL", None, "UNKNOWN_DOCUMENT_ROLE"),
        ("SHIPPING INSTRUCTION\nBILL OF LADING (DRAFT)", "SI", None, "AMBIGUOUS_DOCUMENT_ROLE"),
    ],
)
def test_role_comes_from_content(heading, expected_role, role, code):
    result = extract_document(
        TEXT.replace("SHIPPING INSTRUCTION", heading).encode(),
        "pretend_BL.txt",
        "version",
        expected_role,
    )
    assert result.detected_role == role
    if code:
        assert code in codes(result)
        assert result.needs_review
    else:
        assert not result.needs_review


def test_docx_paragraph_and_table_evidence():
    result = extract_document(docx_bytes(), "source.docx", "docx-v1")
    assert field(result, "shipper").evidence[0].paragraph == 2
    assert field(result, "gross_weight_kg").normalized_value == "123.4"
    assert [e.column for e in field(result, "gross_weight_kg").evidence] == [1, 2]
    assert_evidence(result)


def test_spreadsheet_formulas_are_not_values_and_empty_neighbor_is_not_skipped():
    data = workbook_bytes([["Gross Weight (KG)", "=100+23"], ["Shipper", None, "Not adjacent LTD"]])
    result = extract_document(data, "source.xlsx", "xlsx-v1")
    assert "FORMULA_VALUE" in codes(result)
    assert field(result, "gross_weight_kg").raw_value == "=100+23"
    assert field(result, "gross_weight_kg").normalized_value is None
    assert field(result, "shipper").value_state == "MISSING"
    assert_evidence(result)


@pytest.mark.parametrize(
    ("data", "name", "code"),
    [
        (b"not a pdf", "file.pdf", "FORMAT_MISMATCH"),
        (b"%PDF-1.7\nbroken", "file.pdf", "MALFORMED_DOCUMENT"),
        (b"hello", "file.exe", "UNSUPPORTED_FORMAT"),
        (b"%PDF-1.7", "file.txt", "FORMAT_MISMATCH"),
        (b"PK\x03\x04broken", "file.docx", "MALFORMED_DOCUMENT"),
        (bytes.fromhex("d0cf11e0a1b11ae1"), "file.xlsx", "ENCRYPTED_OR_LEGACY_OFFICE"),
        (b"hello\x00world", "file.txt", "BINARY_TEXT"),
        (b"\xffhello", "file.txt", "UNSUPPORTED_ENCODING"),
        (b"<script>alert(1)</script>", "file.txt", "FORMAT_MISMATCH"),
    ],
)
def test_bad_inputs_return_safe_structured_issues(data, name, code):
    result = extract_document(data, name, "bad-input")
    assert code in codes(result)
    assert result.parsing_status in {"FAILED", "REJECTED"}
    assert all(f.value_state == "UNREADABLE" for f in result.fields)
    assert result.needs_review


def test_pdf_empty_encrypted_and_page_evidence():
    empty = extract_document(pdf_bytes(), "empty.pdf", "empty")
    assert empty.parsing_status == "NO_USABLE_TEXT"
    encrypted = extract_document(pdf_bytes(encrypted=True), "secret.pdf", "secret")
    assert "ENCRYPTED_DOCUMENT" in codes(encrypted)
    parsed = extract_document(pdf_bytes(TEXT, pages=2), "source.pdf", "source")
    assert not parsed.needs_review
    assert {e.page for e in field(parsed, "shipper").evidence} == {1, 2}
    assert_evidence(parsed)


def test_office_type_spoofing_and_unsafe_xml():
    swapped = extract_document(docx_bytes(), "fake.xlsx", "wrong-format")
    assert "FORMAT_MISMATCH" in codes(swapped)
    unsafe = replace_zip_part(
        docx_bytes(),
        "word/document.xml",
        lambda _: b'<!DOCTYPE doc [<!ENTITY x SYSTEM "file:///not-a-real-file">]><doc>&x;</doc>',
    )
    assert "MALFORMED_DOCUMENT" in codes(extract_document(unsafe, "unsafe.docx", "unsafe"))


@pytest.mark.parametrize(
    ("content", "filename", "limits"),
    [
        (TEXT.encode(), "source.txt", {"max_file_bytes": 10}),
        (TEXT.encode(), "source.txt", {"max_source_units": 2}),
        (TEXT.encode(), "source.txt", {"max_text_chars": 10}),
        (pdf_bytes(pages=2), "source.pdf", {"max_pdf_pages": 1}),
        (pdf_bytes(TEXT), "source.pdf", {"max_pdf_stream_bytes": 10}),
        (docx_bytes(), "source.docx", {"max_archive_entries": 1}),
        (docx_bytes(), "source.docx", {"max_archive_bytes": 10}),
        (workbook_bytes(sheets=2), "source.xlsx", {"max_sheets": 1}),
        (workbook_bytes([["a", "b"]]), "source.xlsx", {"max_rows": 1}),
        (workbook_bytes([["a", "b"]]), "source.xlsx", {"max_columns": 1}),
        (workbook_bytes([["a", "b"]]), "source.xlsx", {"max_cells": 2}),
    ],
    ids=[
        "file-bytes",
        "source-units",
        "text-size",
        "pdf-pages",
        "pdf-stream",
        "archive-entries",
        "archive-expansion",
        "sheets",
        "rows",
        "columns",
        "cells",
    ],
)
def test_resource_limits(content, filename, limits):
    result = extract_document(content, filename, "bounded", limits=ExtractionLimits(**limits))
    assert "RESOURCE_LIMIT" in codes(result)
    assert result.parsing_status == "REJECTED"


def test_false_sheet_dimensions_do_not_hide_values_or_bypass_limits():
    data = workbook_bytes([["Gross Weight (KG)", 123]])
    data = replace_zip_part(
        data, "xl/worksheets/sheet1.xml", lambda xml: xml.replace(b"A1:B2", b"A1:A1")
    )
    result = extract_document(data, "source.xlsx", "dimensions")
    assert field(result, "gross_weight_kg").normalized_value == "123"
    bounded = extract_document(
        data, "source.xlsx", "dimensions", limits=ExtractionLimits(max_rows=1)
    )
    assert "RESOURCE_LIMIT" in codes(bounded)


def test_oversized_row_metadata_cannot_bypass_cell_limits():
    data = workbook_bytes([["Gross Weight (KG)", 123]])
    data = replace_zip_part(
        data,
        "xl/worksheets/sheet1.xml",
        lambda xml: xml.replace(b'<row r="2">', b'<row r="99999">'),
    )
    assert "RESOURCE_LIMIT" in codes(extract_document(data, "source.xlsx", "huge-row"))


def test_extreme_container_integer_requires_review_instead_of_crashing():
    result = extract_text(TEXT.replace("6 x 40'HC", "9" * 5000))
    assert field(result, "container_count").value_state == "AMBIGUOUS"
    assert "AMBIGUOUS_CONTAINER_COUNT" in codes(result)


def test_individual_container_table_does_not_establish_total_weight():
    result = extract_text(
        "SHIPPING INSTRUCTION\nCONTAINER NO.\nDESCRIPTION\nGROSS WEIGHT (KG)\n123\n"
    )
    assert field(result, "gross_weight_kg").normalized_value is None
    assert "TOTAL_WEIGHT_REQUIRED" in codes(result)


def test_json_contract_and_filename_independence():
    a = extract_text()
    b = extract_document(TEXT.encode(), "email_160_SI.txt", "source-v1")
    assert a.fields == b.fields
    assert DocumentExtraction.model_validate_json(a.model_dump_json()) == a
    assert all(f.method == "rule" and not f.requires_human_confirmation for f in a.fields)


@pytest.mark.parametrize(("doc_id", "role"), [("", None), ("source", "INVOICE")])
def test_invalid_caller_arguments(doc_id, role):
    with pytest.raises(ValueError):
        extract_document(TEXT.encode(), "source.txt", doc_id, role)


def test_cli_json_and_exit_codes(tmp_path):
    def run(path, *extra):
        return subprocess.run(
            [sys.executable, "-m", "app.extraction", str(path), *extra],
            capture_output=True,
            text=True,
            check=False,
        )

    valid = run(FIXTURES / "email_160_SI.pdf", "--role", "SI")
    assert valid.returncode == 0
    payload = json.loads(valid.stdout)
    assert payload["document_id"].startswith("local-")
    assert not payload["needs_review"]
    missing = run(tmp_path / "missing.pdf")
    assert missing.returncode == 2
    assert json.loads(missing.stdout)["error"]["code"] == "FILE_READ_FAILED"
    corrupt = tmp_path / "bad.pdf"
    corrupt.write_bytes(b"%PDF-broken")
    failed = run(corrupt)
    assert failed.returncode == 2
    assert json.loads(failed.stdout)["parsing_status"] == "FAILED"
    review = run(FIXTURES / "email_055_SI.xlsx", "--document-id", "source-v2")
    assert review.returncode == 0
    assert json.loads(review.stdout)["needs_review"]
    assert json.loads(review.stdout)["document_id"] == "source-v2"
    oversized = tmp_path / "oversized.txt"
    oversized.write_bytes(b"x" * (ExtractionLimits().max_file_bytes + 2))
    rejected = run(oversized)
    assert rejected.returncode == 2
    assert json.loads(rejected.stdout)["error"]["code"] == "RESOURCE_LIMIT"
    assert "content_sha256" not in json.loads(rejected.stdout)
